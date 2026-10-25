"""Writes the "start here" sheet of the Sales X-Ray Videos folder as .xlsx (run through uv with
openpyxl); Drive turns it into a Google Sheet on upload (owner, 25 Oct 2026: one simple sheet, one
row per video, frozen header, brand colour, big readable type, a hidden notes tab):

    uv run --no-project --with openpyxl python sheet_xlsx.py <manifest.json> <links.json> <out.xlsx>

links.json maps each video's path in the folder ("English/01 Join the testers.mp4") to its link.
"""

import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DEEP, MINT, PAPER, INK, WHITE = "0B4A44", "3DEBB0", "F6FAF8", "0C1A18", "FFFFFF"
LANGS = [("en", "English"), ("hi", "Hindi"), ("mr", "Marathi")]
FONT = "Hind"


def clock(s: float) -> str:
    s = round(s)
    return f"{s // 60}:{s % 60:02d}"


def length(v: dict) -> str:
    """The phone video's length: one value, or the range across languages."""
    secs = sorted(x["reel"] for x in v["seconds"].values() if "reel" in x)
    if not secs:
        return ""
    lo, hi = clock(secs[0]), clock(secs[-1])
    return lo if lo == hi else f"{lo}–{hi}"


def link(url: str | None, label: str) -> str:
    return f'=HYPERLINK("{url}","{label}")' if url else ""


def main(manifest_path: str, links_path: str, out: str) -> None:
    m = json.loads(Path(manifest_path).read_text(encoding="utf-8"))
    links = json.loads(Path(links_path).read_text(encoding="utf-8"))
    wb = Workbook()
    ws = wb.active
    ws.title = "Videos"
    cols = ["#", "Video", "What it shows", "Length", "English", "Hindi", "Marathi"]
    widths = [6, 34, 66, 12, 14, 14, 14]
    last = get_column_letter(len(cols))
    deep = PatternFill("solid", fgColor=DEEP)

    ws.merge_cells(f"A1:{last}1")
    ws["A1"] = m["sheet"]
    ws["A1"].font = Font(name=FONT, size=22, bold=True, color=WHITE)
    ws["A1"].fill = deep
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 44
    ws.merge_cells(f"A2:{last}2")
    ws["A2"] = link(m["optin"], "New tester? Join here first: " + m["optin"])
    ws["A2"].font = Font(name=FONT, size=14, bold=True, color=DEEP, underline="single")
    ws["A2"].fill = PatternFill("solid", fgColor=MINT)
    ws["A2"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[2].height = 32
    for i, (name, width) in enumerate(zip(cols, widths, strict=True), 1):
        c = ws.cell(row=3, column=i, value=name)
        c.font = Font(name=FONT, size=14, bold=True, color=WHITE)
        c.fill = deep
        c.alignment = Alignment(horizontal="center" if i != 3 else "left", vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[3].height = 30

    for r, v in enumerate(m["videos"], 4):
        fill = PatternFill("solid", fgColor=PAPER if r % 2 == 0 else WHITE)
        values = [v["number"], v["title"], v["shows"], length(v)]
        values += [link(links.get(v["phone"].get(lang)), "▶ Watch") for lang, _ in LANGS]
        for i, value in enumerate(values, 1):
            c = ws.cell(row=r, column=i, value=value)
            c.fill = fill
            c.font = Font(
                name=FONT,
                size=14,
                bold=i == 2,
                color=DEEP if i >= 5 else INK,
                underline="single" if i >= 5 and value else None,
            )
            c.alignment = Alignment(
                horizontal="left" if i in (2, 3) else "center", vertical="center", wrap_text=i == 3
            )
        ws.row_dimensions[r].height = 42
    ws.freeze_panes = "A4"

    notes = wb.create_sheet("Notes")
    notes.append(["Updated", date.today().isoformat()])
    notes.append(
        ["How to update", "On ac: tutorials.py v2 build <id>, then publish.sh (README, v2)."]
    )
    notes.append([])
    notes.append(["#", "id", "Video", *(f"YouTube (wide) {label}" for _, label in LANGS)])
    for v in m["videos"]:
        wide = [link(links.get(v["wide"].get(lang)), "wide") for lang, _ in LANGS]
        notes.append([v["number"], v["id"], v["title"], *wide])
    notes.sheet_state = "hidden"
    wb.save(out)


if __name__ == "__main__":
    main(*sys.argv[1:4])
