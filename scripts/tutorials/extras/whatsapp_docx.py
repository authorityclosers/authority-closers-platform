"""Word copies of the WhatsApp tester messages (run with uv and python-docx).

uv run --no-project --with python-docx python whatsapp_docx.py <dir with the .txt files>
"""

import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

DEEP = RGBColor(0x0B, 0x4A, 0x44)


def flush(doc: Document, para: list[str]) -> None:
    if para:
        doc.add_paragraph("\n".join(para))
        para.clear()


def convert(txt: Path) -> Path:
    lines = txt.read_text(encoding="utf-8").split("\n")
    doc = Document()
    doc.styles["Normal"].font.name = "Hind"
    doc.styles["Normal"].font.size = Pt(11)
    head = doc.add_heading(lines[0], level=1)
    for run in head.runs:
        run.font.color.rgb = DEEP
    para: list[str] = []
    for line in lines[1:]:
        if line.startswith("=== "):
            flush(doc, para)
            doc.add_heading(line[4:], level=2)
        elif line.startswith("--- "):
            flush(doc, para)
            doc.add_heading(line[4:], level=3)
        elif line.strip():
            para.append(line)
        else:
            flush(doc, para)
    flush(doc, para)
    out = txt.with_suffix(".docx")
    doc.save(out)
    return out


if __name__ == "__main__":
    for txt in sorted(Path(sys.argv[1]).glob("whatsapp-testers-*.txt")):
        print(convert(txt).name)
